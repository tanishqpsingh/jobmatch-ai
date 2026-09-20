import io
import re
from typing import Dict, List, Any
import pymupdf as fitz  # PyMuPDF
import docx
from fastapi import HTTPException, status
from backend.app.schemas.resume import ResumeParseResponse


def extract_text_from_pdf(content: bytes) -> str:
    """Extract full raw text from a PDF document using PyMuPDF."""
    try:
        doc = fitz.open(stream=content, filetype="pdf")
        if doc.page_count == 0:
            raise ValueError("PDF document contains no pages.")

        extracted_text = []
        for page_num in range(doc.page_count):
            page = doc.load_page(page_num)
            text = page.get_text("text")
            if text:
                extracted_text.append(text)
        doc.close()

        full_text = "\n".join(extracted_text).strip()
        if not full_text:
            raise ValueError("No extractable text found in PDF.")

        return full_text
    except Exception as e:
        if isinstance(e, HTTPException):
            raise e
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Malformed or corrupt PDF document: {str(e)}",
        )


def extract_text_from_docx(content: bytes) -> str:
    """Extract full raw text from a DOCX document using python-docx."""
    try:
        doc_stream = io.BytesIO(content)
        document = docx.Document(doc_stream)
        extracted_lines = []

        # Extract text from paragraphs
        for paragraph in document.paragraphs:
            if paragraph.text and paragraph.text.strip():
                extracted_lines.append(paragraph.text.strip())

        # Extract text from tables
        for table in document.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if row_text:
                    extracted_lines.append(" | ".join(row_text))

        full_text = "\n".join(extracted_lines).strip()
        if not full_text:
            raise ValueError("No extractable text found in DOCX document.")

        return full_text
    except Exception as e:
        if isinstance(e, HTTPException):
            raise e
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Malformed or corrupt DOCX document: {str(e)}",
        )


def normalize_text(text: str) -> str:
    """Normalize extracted text by standardizing whitespace and removing control chars."""
    if not text:
        return ""
    # Standardize line endings
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    # Remove non-printable control characters except tab and newline
    normalized = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f-\x9f]', '', normalized)
    # Collapse 3+ consecutive newlines into double newlines
    normalized = re.sub(r'\n{3,}', '\n\n', normalized)
    return normalized.strip()


def extract_contact_info(text: str) -> Dict[str, str | None]:
    """Extract candidate name, email, and phone number deterministically."""
    lines = [line.strip() for line in text.split("\n") if line.strip()]

    # Extract Email
    email_pattern = r'[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
    email_match = re.search(email_pattern, text)
    email = email_match.group(0) if email_match else None

    # Extract Phone
    phone_pattern = r'(?:\+\d{1,3}[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}'
    phone_match = re.search(phone_pattern, text)
    phone = phone_match.group(0) if phone_match else None

    # Extract Name (heuristic: first non-contact header line)
    name = None
    for line in lines[:5]:
        # Skip lines containing email, phone, web links, or standard resume titles
        if "@" in line or "http" in line or re.search(r'\d{5,}', line):
            continue
        if re.search(r'\b(resume|curriculum|vitae|cv|page|contact|profile)\b', line, re.IGNORECASE):
            continue
        # A name line typically has 1-4 words without generic punctuation
        if 1 <= len(line.split()) <= 4 and re.match(r'^[A-Za-z\s\.\'-]+$', line):
            name = line
            break

    return {"name": name, "email": email, "phone": phone}


def extract_sections(text: str) -> Dict[str, List[str]]:
    """Parse text into categorized resume sections deterministically."""
    section_headers = {
        "education": r'\b(education|academic background|qualifications)\b',
        "skills": r'\b(skills|technical skills|core competencies|technologies|expertise)\b',
        "experience": r'\b(experience|work experience|employment history|professional experience)\b',
        "projects": r'\b(projects|personal projects|key projects)\b',
        "certifications": r'\b(certifications|certificates|licenses|training)\b',
    }

    sections: Dict[str, List[str]] = {
        "education": [],
        "skills": [],
        "experience": [],
        "projects": [],
        "certifications": [],
    }

    lines = [line.strip() for line in text.split("\n") if line.strip()]
    current_section = None

    for line in lines:
        line_lower = line.lower()
        matched_section = None

        # Check if line acts as a section header
        for sec_key, pattern in section_headers.items():
            if re.search(pattern, line_lower) and len(line.split()) <= 4:
                matched_section = sec_key
                break

        if matched_section:
            current_section = matched_section
            continue

        if current_section and line:
            # If bullet point or comma-separated list item, split if helpful
            if current_section == "skills" and ("," in line or "•" in line or "|" in line):
                items = re.split(r'[,•|]', line)
                cleaned_items = [item.strip() for item in items if item.strip()]
                sections[current_section].extend(cleaned_items)
            else:
                cleaned_line = line.lstrip("•-* ").strip()
                if cleaned_line:
                    sections[current_section].append(cleaned_line)

    return sections


def parse_resume(filename: str, content: bytes, extension: str) -> ResumeParseResponse:
    """Parse resume content based on extension and return structured schema."""
    if extension == ".pdf":
        raw_text = extract_text_from_pdf(content)
    elif extension == ".docx":
        raw_text = extract_text_from_docx(content)
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file extension: {extension}",
        )

    normalized_text = normalize_text(raw_text)
    contact = extract_contact_info(normalized_text)
    sections = extract_sections(normalized_text)

    return ResumeParseResponse(
        filename=filename,
        name=contact["name"],
        email=contact["email"],
        phone=contact["phone"],
        education=sections["education"],
        skills=sections["skills"],
        experience=sections["experience"],
        projects=sections["projects"],
        certifications=sections["certifications"],
        raw_text=normalized_text,
    )
